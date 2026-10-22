"""Final reports as PDF and PowerPoint, for the Manual and the AI Analyst."""

import io
import json
import re

import pytest
from pptx import Presentation

from app.agent.workflow import BusinessContext, RulesAnalyst, run_workflow
from app.analysis import run_analysis
from app.errors import DataError
from app.profiling import build_profile
from app.report_export import agent_doc, manual_doc, render_pdf, render_pptx
from app.schemas import AnalysisSpec

PPTX_MEDIA = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


@pytest.fixture
def manual_data(sales_ds):
    profile = build_profile(sales_ds, "sales.csv")
    result = run_analysis(
        sales_ds, AnalysisSpec.model_validate({"measures": [{"column": "revenue", "aggregation": "sum"}], "group_by": "region"})
    )
    return {
        "dataset": profile["dataset"],
        "issues": profile["issues"],
        "result": result,
        "measure_key": "sum__revenue",
        "chart_kind": "bar",
    }


@pytest.fixture
def agent_report(sales_ds):
    ctx = BusinessContext()
    return list(run_workflow(sales_ds, "sales.csv", ctx, RulesAnalyst(ctx)))[-1]["report"]


def slides_text(content: bytes) -> list[str]:
    deck = Presentation(io.BytesIO(content))
    out = []
    for slide in deck.slides:
        parts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        parts += [c.text for sh in slide.shapes if sh.has_table for row in sh.table.rows for c in row.cells]
        out.append(" ".join(parts))
    return out


def chart_of(content: bytes):
    deck = Presentation(io.BytesIO(content))
    return next(sh.chart for slide in deck.slides for sh in slide.shapes if sh.has_chart)


def pdf_pages(content: bytes) -> int:
    return len(re.findall(rb"/Type /Page\b", content))


def test_manual_pptx_has_kpis_native_chart_and_quality(manual_data):
    content = render_pptx(manual_doc(manual_data))
    text = slides_text(content)
    assert text[0].startswith("TABULAVIBE AI") and "Sum of revenue by region" in text[0]
    assert "355" in text[1]  # hand-checked total revenue (conftest)
    chart = chart_of(content)
    # Native (editable) PowerPoint chart with the hand-checked group totals, largest first.
    assert list(chart.plots[0].categories) == ["South", "North", "East", "West"]
    assert list(chart.plots[0].series[0].values) == [190, 135, 20, 10]
    assert any("Data health" in t for t in text) and any("Duplicate rows were kept (1 found)" in t for t in text)


def test_manual_pdf_is_a_real_pdf_with_embedded_unicode_font(manual_data):
    content = render_pdf(manual_doc(manual_data))
    assert content.startswith(b"%PDF-") and pdf_pages(content) >= 1
    assert b"DejaVu" in content  # embedded font: renders the rupee sign, arrows and Greek letters


def test_pie_only_when_the_server_allowed_it(manual_data):
    pie = render_pptx(manual_doc({**manual_data, "chart_kind": "pie"}))
    from pptx.enum.chart import XL_CHART_TYPE

    allowed = manual_data["result"]["chart"]["pie_allowed"]
    assert (chart_of(pie).chart_type == XL_CHART_TYPE.DOUGHNUT) is allowed


def test_ai_commentary_is_labelled_when_included(manual_data):
    narrative = {"text": "Revenue is concentrated.\n- South leads\n- West is small", "disclaimer": "AI-generated commentary."}
    text = " ".join(slides_text(render_pptx(manual_doc({**manual_data, "narrative": narrative}))))
    assert "Commentary (AI-generated)" in text and "AI-generated commentary." in text


def test_manual_report_needs_an_analysis():
    with pytest.raises(DataError):
        manual_doc({"dataset": {}, "result": {"kpis": []}})


def test_agent_pptx_has_summary_findings_registers_and_phases(agent_report):
    content = render_pptx(agent_doc(agent_report))
    text = slides_text(content)
    joined = " ".join(text)
    assert "Executive summary" in joined and "Key figures" in joined
    assert "'B' (product) contributes 56.3% of revenue" in joined  # hand-checked: 200 / 355
    insight_slides = [t for t in text if t.lstrip().startswith("Insight F")]
    assert len(insight_slides) == len(agent_report["findings"])
    assert "What might explain it? (not verified)" in insight_slides[0]
    assert "Hypothesis register" in joined and "Workflow phases and gates" in joined
    assert "Reconciled" in joined


def test_agent_pdf_renders(agent_report):
    content = render_pdf(agent_doc(agent_report))
    assert content.startswith(b"%PDF-") and pdf_pages(content) >= 3


def test_unverified_findings_are_flagged_in_the_report(agent_report):
    agent_report["findings"][0]["grounded"] = False
    text = " ".join(slides_text(render_pptx(agent_doc(agent_report))))
    assert "UNVERIFIED figures" in text


def test_agent_report_needs_a_report():
    with pytest.raises(DataError):
        agent_doc({"something": 1})


# ------------------------------------------------------------------ API


def test_export_endpoint_returns_files(client, manual_data, agent_report):
    res = client.post("/api/report/export", json={"format": "pptx", "kind": "manual", "data": manual_data})
    assert res.status_code == 200 and res.headers["content-type"] == PPTX_MEDIA
    assert res.headers["content-disposition"] == 'attachment; filename="tabulavibe_sum_of_revenue_by_region.pptx"'
    res = client.post("/api/report/export", json={"format": "pdf", "kind": "agent", "data": agent_report})
    assert res.status_code == 200 and res.headers["content-type"] == "application/pdf"
    assert res.content.startswith(b"%PDF-")


def test_lone_surrogates_in_the_request_do_not_break_the_files(client, agent_report):
    agent_report["findings"][0]["title"] = "Region \udc81 leads"  # e.g. text mis-decoded by a client
    for fmt in ("pptx", "pdf"):
        body = json.dumps({"format": fmt, "kind": "agent", "data": agent_report})
        res = client.post("/api/report/export", content=body, headers={"Content-Type": "application/json"})
        assert res.status_code == 200, res.text
    # In the typed title the same text is a clean validation error, never a server error.
    body = json.dumps({"format": "pdf", "kind": "agent", "title": "Q\udc81", "data": agent_report})
    assert client.post("/api/report/export", content=body, headers={"Content-Type": "application/json"}).status_code == 422


def test_export_endpoint_validates_input(client, manual_data):
    assert client.post("/api/report/export", json={"format": "docx", "kind": "manual", "data": manual_data}).status_code == 422
    res = client.post("/api/report/export", json={"format": "pdf", "kind": "manual", "data": {}})
    assert res.status_code == 422 and "Run an analysis" in res.json()["error"]["message"]


def test_export_requires_sign_in_when_enabled(manual_data):
    import httpx
    from fastapi.testclient import TestClient

    from app.auth import TokenVerifier, get_token_verifier
    from app.config import Settings
    from app.main import create_app

    from .test_auth import ANON_KEY, FakeAuth

    settings = Settings(
        app_env="test", supabase_url="https://demo.supabase.co", supabase_anon_key=ANON_KEY, require_sign_in=True, _env_file=None
    )
    app = create_app(settings)
    verifier = TokenVerifier(settings.supabase_url, ANON_KEY, 5, transport=httpx.MockTransport(FakeAuth().handler))
    app.dependency_overrides[get_token_verifier] = lambda: verifier
    client = TestClient(app)
    body = {"format": "pdf", "kind": "manual", "data": manual_data}
    assert client.post("/api/report/export", json=body).status_code == 401
    assert client.post("/api/report/export", json=body, headers={"Authorization": "Bearer good-token"}).status_code == 200
