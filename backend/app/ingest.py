"""Safe, in-memory loading of CSV and XLSX uploads into a table of strings.

Rules:
- The file extension must be .csv or .xlsx AND the content must match it.
- Macro-enabled workbooks are rejected; formulas are never evaluated (cached values only).
- Size, row, column and cell limits are enforced while reading.
- Every cell is converted to a trimmed string; type inference happens later.
"""

from __future__ import annotations

import csv
import io
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import PurePath

import pandas as pd

from .config import Settings
from .errors import DataError

SUPPORTED_EXTENSIONS = {".csv", ".xlsx"}
_EXTENSION_HINTS = {
    ".xls": "Legacy .xls files are not supported. Save the file as .xlsx or .csv and try again.",
    ".xlsm": "Macro-enabled workbooks (.xlsm) are not accepted. Save a copy as .xlsx without macros.",
    ".xlsb": "Binary workbooks (.xlsb) are not supported. Save the file as .xlsx or .csv.",
    ".numbers": "Apple Numbers files are not supported. Export to .csv or .xlsx first.",
    ".ods": "OpenDocument spreadsheets are not supported. Save the file as .xlsx or .csv.",
}
_MAX_XLSX_UNCOMPRESSED = 120 * 1024 * 1024  # zip-bomb guard
_MAX_XLSX_ENTRIES = 5_000
_MAX_HEADER_LENGTH = 120
_MAX_CONSECUTIVE_EMPTY_XLSX_ROWS = 10_000


@dataclass
class RawTable:
    frame: pd.DataFrame  # every cell is a str; empty cells are ""
    file_type: str
    sheet_name: str | None = None
    notes: list[str] = field(default_factory=list)
    empty_rows_dropped: int = 0


def load_table(filename: str | None, content: bytes, settings: Settings) -> RawTable:
    ext = PurePath(filename or "").suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        hint = _EXTENSION_HINTS.get(ext, "Only .csv and .xlsx files are supported.")
        raise DataError("unsupported_file_type", hint, 415)
    if not content:
        raise DataError("empty_file", "The file is empty.")
    if len(content) > settings.max_upload_bytes:
        raise DataError(
            "file_too_large",
            f"The file is larger than the {settings.max_upload_mb:g} MB limit.",
            413,
        )

    notes: list[str] = []
    sheet_name = None
    if ext == ".csv":
        rows = _read_csv(content, settings, notes)
        file_type = "csv"
    else:
        if len(content) > settings.max_xlsx_bytes:
            limit = min(settings.max_xlsx_mb, settings.max_upload_mb)
            raise DataError(
                "file_too_large",
                f"Excel files are limited to {limit:g} MB because they are slow to read. "
                f"Save the sheet as CSV to upload up to {settings.max_upload_mb:g} MB.",
                413,
            )
        rows, sheet_name = _read_xlsx(content, settings, notes)
        file_type = "xlsx"

    frame, dropped = _build_frame(rows, settings, notes)
    return RawTable(frame=frame, file_type=file_type, sheet_name=sheet_name, notes=notes, empty_rows_dropped=dropped)


# --------------------------------------------------------------------------- CSV


def _decode(content: bytes, notes: list[str]) -> str:
    if content[:4] == b"PK\x03\x04":
        raise DataError("content_mismatch", "This file looks like a spreadsheet or archive, not a CSV file.", 415)
    if b"\x00" in content[:8192]:
        if content[:2] in (b"\xff\xfe", b"\xfe\xff"):
            try:
                return content.decode("utf-16")
            except UnicodeDecodeError:
                pass
        raise DataError("content_mismatch", "This file contains binary data and is not a valid CSV file.", 415)
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError:
        notes.append("The file is not UTF-8; it was read using the Windows-1252 encoding.")
        try:
            return content.decode("cp1252")
        except UnicodeDecodeError:
            return content.decode("latin-1")


def _read_csv(content: bytes, settings: Settings, notes: list[str]) -> list[list[str]]:
    text = _decode(content, notes)
    if not text.strip():
        raise DataError("empty_file", "The file is empty.")
    sample = text[:65536]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","
    if delimiter != ",":
        names = {";": "semicolon", "\t": "tab", "|": "pipe"}
        notes.append(f"Detected {names.get(delimiter, repr(delimiter))}-separated values.")

    rows: list[list[str]] = []
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, quotechar='"', strict=False)
    try:
        for row in reader:
            rows.append(row)
            if len(rows) > settings.max_rows + 1:
                raise DataError(
                    "too_many_rows",
                    f"The file has more than {settings.max_rows:,} data rows.",
                    413,
                )
    except csv.Error as exc:
        raise DataError("malformed_csv", f"The CSV file could not be parsed (line {reader.line_num}).") from exc
    return rows


# -------------------------------------------------------------------------- XLSX


def _cell_to_str(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, datetime):
        if value.time() == time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, float):
        if value != value:  # NaN
            return ""
        return str(int(value)) if value.is_integer() and abs(value) < 1e15 else repr(value)
    return str(value)


def _read_xlsx(content: bytes, settings: Settings, notes: list[str]) -> tuple[list[list[str]], str]:
    buffer = io.BytesIO(content)
    if content[:4] != b"PK\x03\x04" or not zipfile.is_zipfile(buffer):
        raise DataError("content_mismatch", "This file does not contain a valid .xlsx workbook.", 415)
    try:
        with zipfile.ZipFile(buffer) as archive:
            infos = archive.infolist()
            names = {info.filename for info in infos}
    except zipfile.BadZipFile as exc:
        raise DataError("content_mismatch", "This file does not contain a valid .xlsx workbook.", 415) from exc
    if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
        raise DataError("content_mismatch", "This file does not contain a valid .xlsx workbook.", 415)
    if any(name.lower().endswith("vbaproject.bin") for name in names):
        raise DataError("macros_not_allowed", "Workbooks containing macros are not accepted. Save a copy without macros.", 415)
    # Zip-bomb guard: real workbooks expand about 8-9x; allow 12x the Excel limit (at least 120 MB).
    max_uncompressed = max(_MAX_XLSX_UNCOMPRESSED, 12 * settings.max_xlsx_bytes)
    if len(infos) > _MAX_XLSX_ENTRIES or sum(i.file_size for i in infos) > max_uncompressed:
        raise DataError("file_too_large", "The workbook expands to more data than is allowed.", 413)

    from openpyxl import load_workbook

    buffer.seek(0)
    try:
        workbook = load_workbook(buffer, read_only=True, data_only=True, keep_links=False, keep_vba=False)
    except Exception as exc:  # openpyxl raises many different exception types for corrupt files
        raise DataError("malformed_xlsx", "The workbook could not be read. It may be corrupt or password-protected.") from exc

    try:
        sheets = workbook.worksheets
        if not sheets:
            raise DataError("no_data_rows", "The workbook has no worksheets.")
        sheet = sheets[0]
        if len(sheets) > 1:
            notes.append(f"Only the first worksheet ('{sheet.title}') was read; the workbook has {len(sheets)} sheets.")
        notes.append("Formulas are not evaluated; the values last saved by Excel are used.")

        rows: list[list[str]] = []
        non_empty = 0
        consecutive_empty = 0
        for raw_row in sheet.iter_rows(values_only=True):
            row = [_cell_to_str(v) for v in raw_row]
            while row and row[-1] == "":
                row.pop()
            if not row:
                consecutive_empty += 1
                if consecutive_empty > _MAX_CONSECUTIVE_EMPTY_XLSX_ROWS:
                    break
                rows.append(row)
                continue
            consecutive_empty = 0
            if len(row) > settings.max_columns:
                raise DataError("too_many_columns", f"The file has more than {settings.max_columns} columns.", 413)
            rows.append(row)
            non_empty += 1
            if non_empty > settings.max_rows + 1:
                raise DataError("too_many_rows", f"The file has more than {settings.max_rows:,} data rows.", 413)
        return rows, sheet.title
    finally:
        workbook.close()


# ------------------------------------------------------------------------ common


def _clean_headers(raw: list[str], notes: list[str]) -> list[str]:
    headers: list[str] = []
    seen: dict[str, int] = {}
    renamed_blank = renamed_dup = 0
    for index, value in enumerate(raw, start=1):
        name = " ".join(value.split())[:_MAX_HEADER_LENGTH]
        if not name:
            name = f"column_{index}"
            renamed_blank += 1
        key = name.lower()
        if key in seen:
            seen[key] += 1
            name = f"{name}_{seen[key]}"
            renamed_dup += 1
        else:
            seen[key] = 1
        headers.append(name)
    if renamed_blank:
        notes.append(f"{renamed_blank} blank column header(s) were named column_N.")
    if renamed_dup:
        notes.append(f"{renamed_dup} duplicate column header(s) were given a numeric suffix.")
    return headers


def _build_frame(rows: list[list[str]], settings: Settings, notes: list[str]) -> tuple[pd.DataFrame, int]:
    # Skip leading empty rows to find the header row.
    start = 0
    while start < len(rows) and not any(c.strip() for c in rows[start]):
        start += 1
    if start >= len(rows):
        raise DataError("empty_file", "The file does not contain any data.")
    header_raw = list(rows[start])
    body = rows[start + 1 :]

    # Drop trailing empty header cells only if the column is empty everywhere.
    width = max([len(header_raw)] + [len(r) for r in body])
    while width > 0 and (len(header_raw) < width or not header_raw[width - 1].strip()):
        if any(len(r) >= width and r[width - 1].strip() for r in body):
            break
        width -= 1
    if width == 0:
        raise DataError("empty_file", "The file does not contain any column headers.")
    if width > settings.max_columns:
        raise DataError("too_many_columns", f"The file has more than {settings.max_columns} columns.", 413)
    header_raw = (header_raw + [""] * width)[:width]
    headers = _clean_headers(header_raw, notes)

    data: list[list[str]] = []
    dropped = 0
    for row in body:
        cells = [c.strip() for c in row[:width]]
        if not any(cells):
            dropped += 1
            continue
        if len(cells) < width:
            cells.extend([""] * (width - len(cells)))
        data.append(cells)

    if not data:
        raise DataError("no_data_rows", "The file has column headers but no data rows.")
    if len(data) > settings.max_rows:
        raise DataError("too_many_rows", f"The file has more than {settings.max_rows:,} data rows.", 413)
    if len(data) * width > settings.max_cells:
        raise DataError("too_many_cells", f"The file has more than {settings.max_cells:,} cells.", 413)

    # Leading/trailing empty rows are only worth mentioning when they occur between data rows.
    interior_dropped = dropped - _count_trailing_empty(body, width)
    if interior_dropped > 0:
        notes.append(f"{interior_dropped} completely empty row(s) were skipped.")

    frame = pd.DataFrame(data, columns=headers, dtype=object)
    return frame, max(interior_dropped, 0)


def _count_trailing_empty(body: list[list[str]], width: int) -> int:
    count = 0
    for row in reversed(body):
        if any(c.strip() for c in row[:width]):
            break
        count += 1
    return count
