"""Final reports as PDF (A4) and PowerPoint (16:9) files.

Two builders turn an analysis into a neutral `Doc`:
  * `manual_doc`: the Manual Analyst's report (profile summary + analysis result + optional AI commentary)
  * `agent_doc`:  the AI Analyst's report (phases 0-8, findings, validation, evidence)
Two renderers write the same `Doc` as a PDF (ReportLab, vector charts) or a PPTX deck (python-pptx,
native editable charts). Every figure comes from the server's earlier calculations; this module only
formats (Indian digit grouping, as in the app).
"""

from __future__ import annotations

import io
import threading
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, Literal
from xml.sax.saxutils import escape

from . import __version__
from .agent.workflow import fmt, fmt_p
from .errors import DataError

BRAND = (0x4F, 0x46, 0xE5)  # indigo, --accent
BRAND_2 = (0x7C, 0x3A, 0xED)  # violet, --accent-2
INK = (0x0F, 0x12, 0x22)
INK_2 = (0x47, 0x4C, 0x63)
MUTED = (0x67, 0x6C, 0x84)
LINE = (0xE4, 0xE6, 0xEF)
SURFACE_2 = (0xF3, 0xF4, 0xF9)
SERIES = [
    (0x2A, 0x78, 0xD6),
    (0xEB, 0x68, 0x34),
    (0x1B, 0xAF, 0x7A),
    (0xED, 0xA1, 0x00),
    (0xE8, 0x7B, 0xA4),
    (0x00, 0x83, 0x00),
    (0x62, 0x50, 0xD6),
    (0xE3, 0x49, 0x48),
]
CONFIDENCE_COLOUR = {"high": (0x04, 0x78, 0x57), "medium": BRAND, "low": (0x8A, 0x4B, 0x06)}
FONTS_DIR = Path(__file__).parent / "fonts"  # DejaVu Sans (see fonts/LICENSE_DEJAVU): covers ₹, →, η, ρ, accents
_fonts_lock = threading.Lock()
MAX_TEXT = 1200
MAX_CHART_POINTS = 24
MAX_TABLE_ROWS = 40


# ====================================================================== document model


@dataclass
class Kpis:
    items: list[tuple[str, str, str]]  # label, value, note


@dataclass
class Chart:
    kind: Literal["bar", "line", "pie"]
    title: str
    categories: list[str]
    values: list[float | None]
    series_name: str
    note: str = ""


@dataclass
class TableBlock:
    columns: list[str]
    rows: list[list[str]]
    numeric: set[int] = field(default_factory=set)
    note: str = ""
    widths: list[float] | None = None  # relative column widths; equal when omitted


@dataclass
class Bullets:
    items: list[str]


@dataclass
class Text:
    paragraphs: list[str]


@dataclass
class FindingCard:
    title: str
    badges: list[str]
    colour: tuple[int, int, int]
    fields: list[tuple[str, str]]


Block = Kpis | Chart | TableBlock | Bullets | Text | FindingCard


@dataclass
class Section:
    title: str
    blocks: list[Block]


@dataclass
class Doc:
    title: str
    subtitle: str
    meta: list[str]
    sections: list[Section]
    footer: str


def _clip(text: Any, limit: int = MAX_TEXT) -> str:
    s = " ".join(str(text or "").split())
    return s if len(s) <= limit else s[: limit - 1] + "…"


def _num(value: Any) -> float | None:
    try:
        out = float(value)
    except (TypeError, ValueError):
        return None
    return out if out == out and abs(out) != float("inf") else None


def _today() -> str:
    return datetime.now(UTC).strftime("%d %B %Y")


# ====================================================================== builders


def manual_doc(data: dict[str, Any], title: str | None = None) -> Doc:
    """Report for the Manual Analyst: KPIs, chart, results table, commentary, scope and data quality."""
    result, dataset = data.get("result"), data.get("dataset")
    if not isinstance(result, dict) or not isinstance(dataset, dict) or not result.get("kpis"):
        raise DataError("invalid_request", "Run an analysis before downloading the report.")
    measures = result.get("measures") or []
    measure_key = data.get("measure_key") or (measures[0]["key"] if measures else "")
    measure = next((m for m in measures if m.get("key") == measure_key), measures[0] if measures else {"key": "", "label": "Value"})
    grouping = result.get("grouping")
    by = (f"{grouping['time_grain']} ({grouping['column']})" if grouping.get("time_grain") else grouping["column"]) if grouping else None
    heading = _clip(title or (f"{measure['label']} by {by}" if by else measure["label"]), 120)
    sections: list[Section] = []

    sections.append(
        Section(
            "Key figures",
            [Kpis([(_clip(k.get("label"), 60), fmt(_num(k.get("value"))), _clip(k.get("definition"), 220)) for k in result["kpis"][:6]])],
        )
    )

    rows = result.get("rows") or []
    if grouping and rows:
        kind: Literal["bar", "line", "pie"] = "line" if grouping.get("time_grain") else "bar"
        if data.get("chart_kind") == "pie" and (result.get("chart") or {}).get("pie_allowed"):
            kind = "pie"
        shown = rows if kind == "line" else rows[:MAX_CHART_POINTS]
        sections.append(
            Section(
                heading,
                [
                    Chart(
                        kind,
                        heading,
                        [_clip(r.get("group"), 40) for r in shown],
                        [_num((r.get("values") or {}).get(measure["key"])) for r in shown],
                        measure["label"],
                        "Top groups shown." if len(rows) > len(shown) else "",
                    )
                ],
            )
        )
        cols = [grouping["column"], "Rows", *[m.get("label", "") for m in measures]]
        table_rows = [
            [
                _clip(r.get("group"), 50),
                fmt(_num(r.get("row_count")), 0),
                *[fmt(_num((r.get("values") or {}).get(m.get("key")))) for m in measures],
            ]
            for r in rows[:MAX_TABLE_ROWS]
        ]
        note = f"First {MAX_TABLE_ROWS} of {len(rows)} groups." if len(rows) > MAX_TABLE_ROWS else ""
        sections.append(Section("Detailed results", [TableBlock(cols, table_rows, set(range(1, len(cols))), note)]))

    narrative = data.get("narrative")
    if isinstance(narrative, dict) and narrative.get("text"):
        lines = [ln.strip("-• ").strip() for ln in str(narrative["text"]).splitlines() if ln.strip()]
        sections.append(
            Section(
                "Commentary (AI-generated)",
                [
                    Text([_clip(lines[0], 600)]),
                    Bullets([_clip(ln, 300) for ln in lines[1:8]]),
                    Text([_clip(narrative.get("disclaimer"), 300)]),
                ],
            )
        )

    spec = result.get("spec") or {}
    scope = [
        f"{fmt(_num(result.get('rows_analyzed')), 0)} of {fmt(_num(dataset.get('row_count')), 0)} rows analysed across {dataset.get('column_count')} columns."
    ]
    scope.append(
        "Exact duplicate rows were excluded."
        if spec.get("remove_duplicates")
        else f"Duplicate rows were kept ({dataset.get('duplicate_row_count', 0)} found)."
    )
    if spec.get("merge_label_variants"):
        scope.append("Labels differing only by letter case or spacing were merged.")
    for f in spec.get("filters") or []:
        scope.append(
            f"Filter on {f.get('column')}: "
            + (", ".join(map(str, f["values"][:12])) if f.get("values") else f"{f.get('start') or 'start'} to {f.get('end') or 'end'}")
        )
    scope += [_clip(n, 300) for n in (result.get("notes") or [])[:6]]
    sections.append(
        Section(
            "Scope and definitions",
            [Bullets(scope), Bullets([f"{_clip(k.get('label'), 60)}: {_clip(k.get('definition'), 240)}" for k in result["kpis"][:6]])],
        )
    )

    issues = [i for i in (data.get("issues") or []) if isinstance(i, dict) and i.get("severity") == "warning"]
    quality = [
        f"Data health {fmt(_num(dataset.get('valid_cell_pct')), 1)}% (cells present and valid).",
        f"{fmt(_num(dataset.get('missing_cell_count')), 0)} missing cells ({fmt(_num(dataset.get('missing_cell_pct')), 2)}%) and {fmt(_num(dataset.get('duplicate_row_count')), 0)} duplicate rows detected.",
        "Missing and invalid values are excluded from calculations and were not altered.",
        *[_clip(i.get("message"), 240) for i in issues[:8]],
    ]
    sections.append(Section("Data quality", [Bullets(quality)]))

    source = _clip(dataset.get("file_name") or "uploaded file", 80) + (
        f" (sheet {dataset['sheet_name']})" if dataset.get("sheet_name") else ""
    )
    return Doc(
        title=heading,
        subtitle="Business report · Manual Analyst",
        meta=[f"Source: {source}", f"Prepared {_today()}", f"Rows analysed: {fmt(_num(result.get('rows_analyzed')), 0)}"],
        sections=sections,
        footer=f"TabulaVibe AI {__version__} · Figures were calculated from the uploaded file; review before relying on them for decisions.",
    )


def agent_doc(report: dict[str, Any]) -> Doc:
    """Report for the AI Analyst: summary, KPIs, charts, insight cards, quality, method, limitations."""
    if not isinstance(report, dict) or not isinstance(report.get("findings"), list) or "dashboard" not in report:
        raise DataError("invalid_request", "Run the AI analyst before downloading its report.")
    evidence = {e.get("id"): e for e in report.get("evidence") or [] if isinstance(e, dict)}
    dataset = report.get("dataset") or {}
    findings = [f for f in report["findings"] if isinstance(f, dict)]
    grounded = [f for f in findings if f.get("grounded")]
    sections: list[Section] = []

    summary = [f"{_clip(f.get('title'), 160)}. Action: {_clip(f.get('recommended_action'), 200)}" for f in grounded[:6]] or [
        "No finding met the evidence threshold."
    ]
    if report.get("ai_note"):
        summary.append(_clip(report["ai_note"], 300))
    sections.append(Section("Executive summary", [Bullets(summary)]))

    kpis = [
        (_clip(k.get("label"), 60), fmt(_num(k.get("value"))) + ("%" if k.get("unit") == "%" else ""), _clip(k.get("definition"), 200))
        for k in (report["dashboard"].get("kpis") or [])[:6]
    ]
    if kpis:
        sections.append(Section("Key figures", [Kpis(kpis)]))

    charts = report["dashboard"].get("charts") or []
    trend = next((evidence.get(c.get("evidence_id")) for c in charts if c.get("type") == "trend"), None)
    if trend and trend.get("result", {}).get("periods"):
        r = trend["result"]
        name = f"{'Average' if r.get('aggregation') == 'mean' else 'Total'} {r.get('measure') or 'rows'}"
        note = f"Partial periods (excluded from comparisons): {', '.join(r['partial_periods'])}." if r.get("partial_periods") else ""
        sections.append(
            Section(
                f"{name} by {r.get('grain', 'period')}",
                [
                    Chart(
                        "line",
                        f"{name} by {r.get('grain')}",
                        [p["period"] for p in r["periods"]],
                        [_num(p.get("value")) for p in r["periods"]],
                        name,
                        note,
                    )
                ],
            )
        )
    segment = next((evidence.get(c.get("evidence_id")) for c in charts if c.get("type") == "bar"), None)
    if segment and segment.get("result", {}).get("segments"):
        r = segment["result"]
        segs = r["segments"][:MAX_CHART_POINTS]
        sections.append(
            Section(
                f"{r.get('measure')} by {r.get('dimension')}",
                [
                    Chart(
                        "bar",
                        f"Total {r.get('measure')} by {r.get('dimension')}",
                        [_clip(s.get("segment"), 40) for s in segs],
                        [_num(s.get("total")) for s in segs],
                        f"Total {r.get('measure')}",
                        f"Top segment {fmt((_num(r.get('top_share')) or 0) * 100, 1)}% · top 3 {fmt((_num(r.get('top_3_share')) or 0) * 100, 1)}%",
                    )
                ],
            )
        )

    for f in findings[:10]:
        badges = [
            f"{f.get('confidence', 'low')} confidence",
            {"observed_fact": "Observed fact", "statistical_association": "Statistical association", "data_quality": "Data quality"}.get(
                f.get("kind"), "Finding"
            ),
        ]
        if not f.get("grounded"):
            badges.append("UNVERIFIED figures")
        fields = [
            ("What happened?", f.get("what_happened")),
            ("Where?", f.get("where")),
            ("How large?", f.get("magnitude")),
            ("Compared with", " → ".join(x for x in [f.get("baseline"), f.get("comparison")] if x)),
            ("What might explain it? (not verified)", "; ".join(f.get("possible_explanations") or [])),
            ("Action to consider", f.get("recommended_action")),
            ("How to measure success", f.get("success_metric")),
            ("Limitations", f.get("limitations")),
            ("Evidence", ", ".join(f.get("evidence_ids") or [])),
        ]
        sections.append(
            Section(
                f"Insight {f.get('id', '')}",
                [
                    FindingCard(
                        _clip(f.get("title"), 160),
                        badges,
                        CONFIDENCE_COLOUR.get(f.get("confidence"), MUTED),
                        [(k, _clip(v, 420)) for k, v in fields if v],
                    )
                ],
            )
        )

    du = report.get("data_understanding") or {}
    quality = [f"Grain: {_clip(du.get('grain'), 160)}. {_clip(du.get('grain_basis'), 200)}"]
    quality += [
        f"[{i.get('severity')}] {_clip(i.get('issue'), 100)}: {_clip(i.get('detail'), 220)}"
        for i in (du.get("quality_issues") or [])[:8]
        if isinstance(i, dict)
    ]
    validation = report.get("validation") or {}
    recon = [
        [
            f"{'Average' if r.get('aggregation') == 'mean' else 'Total'} {r.get('measure')}",
            fmt(_num(r.get("engine_value"))),
            fmt(_num(r.get("independent_value"))),
            "Reconciled" if r.get("reconciled") else "MISMATCH",
        ]
        for r in validation.get("reconciliation") or []
    ]
    tests = (validation.get("multiple_testing") or {}).get("tests") or []
    blocks: list[Block] = [Bullets(quality)]
    if recon:
        blocks.append(
            TableBlock(
                ["Measure", "Engine", "Independent", "Status"],
                recon,
                {1, 2},
                f"Holm-Bonferroni correction across {len(tests)} statistical test(s), alpha 0.05.",
            )
        )
    sections.append(Section("Data quality and validation", blocks))

    charter = report.get("charter") or {}
    plan = [f"Business context: {_clip(charter.get('business_context'), 200)}", f"Audience: {_clip(charter.get('stakeholders'), 120)}"]
    plan += [f"Objective: {_clip(o, 200)}" for o in (charter.get("objectives") or [])[:3]]
    plan += [f"Assumption: {_clip(a, 200)}" for a in (charter.get("assumptions") or [])[:5]]
    questions = [
        [q.get("id", ""), _clip(q.get("question"), 140), q.get("priority", ""), ", ".join(q.get("evidence_ids") or []) or "none"]
        for q in (report.get("questions") or [])[:12]
    ]
    hypotheses = [
        [
            h.get("id", ""),
            _clip(h.get("statement"), 140),
            str(h.get("result", "")).replace("_", " "),
            fmt_p(_num(h.get("adjusted_p_value"))),
        ]
        for h in (report.get("hypotheses") or [])[:12]
    ]
    sections.append(Section("Plan and method", [Bullets(plan)]))
    if questions:
        sections.append(
            Section(
                "Question register", [TableBlock(["ID", "Question", "Priority", "Evidence"], questions, widths=[0.07, 0.63, 0.12, 0.18])]
            )
        )
    if hypotheses:
        sections.append(
            Section(
                "Hypothesis register",
                [TableBlock(["ID", "Hypothesis", "Result", "Adjusted p"], hypotheses, {3}, widths=[0.07, 0.63, 0.15, 0.15])],
            )
        )
    phases = [[str(p.get("id")), _clip(p.get("name"), 60), str(p.get("status", "")).replace("_", " ")] for p in report.get("phases") or []]
    sections.append(
        Section(
            "Limitations and next steps",
            [Bullets([_clip(x, 300) for x in (report.get("limitations") or []) + (report.get("outstanding_questions") or [])][:10])],
        )
    )
    sections.append(Section("Workflow phases and gates", [TableBlock(["Phase", "Name", "Status"], phases, widths=[0.12, 0.58, 0.30])]))
    mode = f"AI analyst ({report.get('model')})" if report.get("mode") == "ai" else "Rules-based analyst"
    return Doc(
        title=f"Analysis of {_clip(dataset.get('file_name') or 'uploaded file', 80)}",
        subtitle=f"Business report · AI Analyst · {mode}",
        meta=[
            f"Prepared {_today()}",
            f"{fmt(_num(dataset.get('rows')), 0)} rows · {dataset.get('columns')} columns",
            f"{len(report.get('evidence') or [])} calculations recorded as evidence",
        ],
        sections=sections,
        footer=f"TabulaVibe AI {__version__} · Findings describe associations in this file, not causes. Every figure was computed in Python.",
    )


# ====================================================================== PowerPoint


def render_pptx(doc: Doc) -> bytes:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Inches, Pt

    def rgb(c: tuple[int, int, int]) -> RGBColor:
        return RGBColor(*c)

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.core_properties.title = doc.title
    prs.core_properties.author = "TabulaVibe AI"
    blank = prs.slide_layouts[6]
    width = prs.slide_width
    margin = Inches(0.6)
    content_w = width - 2 * margin

    def text(slide, x, y, w, h, value: str, size: int, colour=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
        box = slide.shapes.add_textbox(x, y, w, h)
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.05)
        p = tf.paragraphs[0]
        p.alignment = align
        run = p.add_run()
        run.text = value
        run.font.size, run.font.bold, run.font.color.rgb = Pt(size), bold, rgb(colour)
        return box

    def new_slide(title: str):
        slide = prs.slides.add_slide(blank)
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, width, Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = rgb(BRAND)
        bar.line.fill.background()
        text(slide, margin, Inches(0.35), content_w, Inches(0.8), title, 26, INK, bold=True)
        text(slide, margin, prs.slide_height - Inches(0.45), content_w, Inches(0.3), doc.footer, 9, MUTED)
        return slide

    # cover
    slide = prs.slides.add_slide(blank)
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, width, prs.slide_height)
    band.fill.gradient()
    band.fill.gradient_angle = 135
    band.fill.gradient_stops[0].color.rgb = rgb(BRAND)
    band.fill.gradient_stops[1].color.rgb = rgb(BRAND_2)
    band.line.fill.background()
    text(slide, margin, Inches(1.2), content_w, Inches(0.5), "TABULAVIBE AI", 14, (0xE0, 0xE7, 0xFF), bold=True)
    text(slide, margin, Inches(2.0), content_w, Inches(2.0), doc.title, 40, (0xFF, 0xFF, 0xFF), bold=True)
    text(slide, margin, Inches(4.1), content_w, Inches(0.6), doc.subtitle, 18, (0xE0, 0xE7, 0xFF))
    text(slide, margin, Inches(5.0), content_w, Inches(1.2), "   ·   ".join(doc.meta), 14, (0xFF, 0xFF, 0xFF))

    for section in doc.sections:
        blocks = list(section.blocks)
        if any(isinstance(b, FindingCard) for b in blocks):
            for card in [b for b in blocks if isinstance(b, FindingCard)]:
                slide = new_slide(section.title)
                text(slide, margin, Inches(1.15), content_w, Inches(0.8), card.title, 20, INK, bold=True)
                text(slide, margin, Inches(1.85), content_w, Inches(0.35), "  ·  ".join(card.badges), 12, card.colour, bold=True)
                col_w = (content_w - Inches(0.4)) // 2
                fields = card.fields
                half = (len(fields) + 1) // 2
                for col, chunk in enumerate((fields[:half], fields[half:])):
                    box = slide.shapes.add_textbox(margin + col * (col_w + Inches(0.4)), Inches(2.35), col_w, Inches(4.6))
                    tf = box.text_frame
                    tf.word_wrap = True
                    first = True
                    for label, value in chunk:
                        p = tf.paragraphs[0] if first else tf.add_paragraph()
                        first = False
                        r = p.add_run()
                        r.text = label
                        r.font.size, r.font.bold, r.font.color.rgb = Pt(11), True, rgb(MUTED)
                        p2 = tf.add_paragraph()
                        r2 = p2.add_run()
                        r2.text = value
                        r2.font.size, r2.font.color.rgb = Pt(13), rgb(INK_2)
                        p2.space_after = Pt(8)
            continue
        slide = new_slide(section.title)
        y = Inches(1.3)
        bottom = prs.slide_height - Inches(0.6)
        for block in blocks:
            if isinstance(block, Kpis):
                per_row = 3
                card_w = (content_w - Inches(0.3) * (per_row - 1)) // per_row
                card_h = Inches(1.9)
                for i, (label, value, note) in enumerate(block.items):
                    x = margin + (i % per_row) * (card_w + Inches(0.3))
                    top = y + (i // per_row) * (card_h + Inches(0.25))
                    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, top, card_w, card_h)
                    card.adjustments[0] = 0.08
                    card.fill.solid()
                    card.fill.fore_color.rgb = rgb((0xFF, 0xFF, 0xFF))
                    card.line.color.rgb = rgb(LINE)
                    text(slide, x + Inches(0.2), top + Inches(0.15), card_w - Inches(0.4), Inches(0.35), label, 12, MUTED, bold=True)
                    text(slide, x + Inches(0.2), top + Inches(0.5), card_w - Inches(0.4), Inches(0.6), value, 28, INK, bold=True)
                    text(slide, x + Inches(0.2), top + Inches(1.15), card_w - Inches(0.4), Inches(0.7), note, 9, MUTED)
                y += ((len(block.items) + per_row - 1) // per_row) * (card_h + Inches(0.25))
            elif isinstance(block, Chart):
                data = CategoryChartData()
                data.categories = block.categories or ["(none)"]
                data.add_series(block.series_name, [v if v is not None else None for v in block.values] or [0])
                kind = {"bar": XL_CHART_TYPE.BAR_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS, "pie": XL_CHART_TYPE.DOUGHNUT}[block.kind]
                height = max(Inches(2.5), bottom - y - (Inches(0.4) if block.note else 0))
                chart = slide.shapes.add_chart(kind, margin, y, content_w, height, data).chart
                chart.font.size = Pt(11)
                chart.font.color.rgb = rgb(INK_2)
                indian = r"[>=10000000]##\,##\,##\,##0;[>=100000]##\,##\,##0;##,##0"
                if block.kind == "pie":
                    chart.has_legend = True
                    chart.legend.position = XL_LEGEND_POSITION.RIGHT
                    chart.legend.include_in_layout = False
                    for i, point in enumerate(chart.plots[0].series[0].points):
                        point.format.fill.solid()
                        point.format.fill.fore_color.rgb = rgb(SERIES[i % len(SERIES)])
                else:
                    chart.has_legend = False
                    plot = chart.plots[0]
                    series = plot.series[0]
                    if block.kind == "bar":
                        series.format.fill.solid()
                        series.format.fill.fore_color.rgb = rgb(SERIES[0])
                        plot.gap_width = 60
                        plot.has_data_labels = True
                        plot.data_labels.number_format = indian
                        plot.data_labels.number_format_is_linked = False
                        plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
                        chart.category_axis.reverse_order = True  # largest group on top, like the app
                    else:
                        series.format.line.color.rgb = rgb(SERIES[0])
                        series.format.line.width = Pt(2.25)
                        series.smooth = False
                    chart.value_axis.tick_labels.number_format = indian
                    chart.value_axis.tick_labels.number_format_is_linked = False
                    chart.value_axis.major_gridlines.format.line.color.rgb = rgb(LINE)
                    chart.value_axis.format.line.fill.background()
                y += height
                if block.note:
                    text(slide, margin, y, content_w, Inches(0.35), block.note, 10, MUTED)
                    y += Inches(0.4)
            elif isinstance(block, TableBlock):
                rows = block.rows[:12]
                n_rows = len(rows) + 1
                row_h = Inches(0.36)
                shape = slide.shapes.add_table(n_rows, len(block.columns), margin, y, content_w, Emu(row_h * n_rows))
                table = shape.table
                if block.widths:
                    total = sum(block.widths)
                    for c, share in enumerate(block.widths):
                        table.columns[c].width = Emu(int(content_w * share / total))
                for c, name in enumerate(block.columns):
                    cell = table.cell(0, c)
                    cell.text = name
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = rgb(SURFACE_2)
                    para = cell.text_frame.paragraphs[0]
                    para.runs[0].font.size, para.runs[0].font.bold, para.runs[0].font.color.rgb = Pt(11), True, rgb(INK_2)
                    para.alignment = PP_ALIGN.RIGHT if c in block.numeric else PP_ALIGN.LEFT
                for r, row in enumerate(rows, start=1):
                    for c, value in enumerate(row):
                        cell = table.cell(r, c)
                        cell.text = value
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = rgb((0xFF, 0xFF, 0xFF))
                        para = cell.text_frame.paragraphs[0]
                        para.runs[0].font.size, para.runs[0].font.color.rgb = Pt(11), rgb(INK)
                        para.alignment = PP_ALIGN.RIGHT if c in block.numeric else PP_ALIGN.LEFT
                y += Emu(row_h * n_rows) + Inches(0.15)
                note = block.note or (f"First 12 of {len(block.rows)} rows; the PDF report lists more." if len(block.rows) > 12 else "")
                if note:
                    text(slide, margin, y, content_w, Inches(0.35), note, 10, MUTED)
                    y += Inches(0.4)
            elif isinstance(block, Bullets | Text):
                items = block.items if isinstance(block, Bullets) else block.paragraphs
                if not items:
                    continue
                box = slide.shapes.add_textbox(margin, y, content_w, max(Inches(0.5), bottom - y))
                tf = box.text_frame
                tf.word_wrap = True
                size = 16 if len(items) <= 5 else 13
                for i, item in enumerate(items[:12]):
                    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    run = p.add_run()
                    run.text = ("•  " if isinstance(block, Bullets) else "") + item
                    run.font.size, run.font.color.rgb = Pt(size), rgb(INK_2)
                    p.space_after = Pt(6)
                y += Inches(0.42) * min(len(items), 12) + Inches(0.3)
            if y > bottom:
                break
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ====================================================================== PDF


def _register_pdf_fonts() -> None:
    from reportlab.lib.fonts import addMapping
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    with _fonts_lock:
        if "DejaVu" in pdfmetrics.getRegisteredFontNames():
            return
        pdfmetrics.registerFont(TTFont("DejaVu", str(FONTS_DIR / "DejaVuSans.ttf")))
        pdfmetrics.registerFont(TTFont("DejaVu-Bold", str(FONTS_DIR / "DejaVuSans-Bold.ttf")))
        addMapping("DejaVu", 0, 0, "DejaVu")
        addMapping("DejaVu", 1, 0, "DejaVu-Bold")


def render_pdf(doc: Doc) -> bytes:
    from reportlab.graphics.charts.barcharts import HorizontalBarChart
    from reportlab.graphics.charts.linecharts import HorizontalLineChart
    from reportlab.graphics.charts.piecharts import Pie
    from reportlab.graphics.shapes import Drawing
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    _register_pdf_fonts()

    def col(c: tuple[int, int, int]):
        return colors.Color(c[0] / 255, c[1] / 255, c[2] / 255)

    styles = {
        "h1": ParagraphStyle("h1", fontName="DejaVu-Bold", fontSize=22, leading=27, textColor=col(INK), spaceAfter=4),
        "sub": ParagraphStyle("sub", fontName="DejaVu", fontSize=11, leading=15, textColor=col(BRAND)),
        "meta": ParagraphStyle("meta", fontName="DejaVu", fontSize=9, leading=13, textColor=col(MUTED)),
        "h2": ParagraphStyle(
            "h2", fontName="DejaVu-Bold", fontSize=14, leading=18, textColor=col(INK), spaceBefore=12, spaceAfter=6, keepWithNext=1
        ),
        "body": ParagraphStyle("body", fontName="DejaVu", fontSize=10, leading=14, textColor=col(INK_2)),
        "small": ParagraphStyle("small", fontName="DejaVu", fontSize=8, leading=11, textColor=col(MUTED)),
        "label": ParagraphStyle("label", fontName="DejaVu-Bold", fontSize=8, leading=11, textColor=col(MUTED)),
        "kpi": ParagraphStyle("kpi", fontName="DejaVu-Bold", fontSize=17, leading=21, textColor=col(INK)),
        "cell": ParagraphStyle("cell", fontName="DejaVu", fontSize=8.5, leading=11, textColor=col(INK)),
        "cellr": ParagraphStyle("cellr", fontName="DejaVu", fontSize=8.5, leading=11, textColor=col(INK), alignment=TA_RIGHT),
        "head": ParagraphStyle("head", fontName="DejaVu-Bold", fontSize=8.5, leading=11, textColor=col(INK_2)),
        "card": ParagraphStyle("card", fontName="DejaVu-Bold", fontSize=12, leading=16, textColor=col(INK)),
    }

    def p(value: str, style: str) -> Paragraph:
        return Paragraph(escape(value), styles[style])

    page_w = A4[0] - 36 * mm
    story: list[Any] = [p(doc.subtitle.upper(), "sub"), Spacer(1, 4), p(doc.title, "h1"), p("   ·   ".join(doc.meta), "meta"), Spacer(1, 6)]
    rule = Table([[""]], colWidths=[page_w], rowHeights=[2])
    rule.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), col(BRAND))]))
    story.append(rule)

    def number_label(v: float) -> str:
        return fmt(v, 0) if abs(v) >= 100 else fmt(v)

    for section in doc.sections:
        story.append(p(section.title, "h2"))
        for block in section.blocks:
            if isinstance(block, Kpis):
                cells = [[p(label, "label"), p(value, "kpi"), p(note, "small")] for label, value, note in block.items]
                grid = [cells[i : i + 3] for i in range(0, len(cells), 3)]
                grid = [row + [""] * (3 - len(row)) for row in grid]
                table = Table([[Table([[c[0]], [c[1]], [c[2]]]) if c else "" for c in row] for row in grid], colWidths=[page_w / 3] * 3)
                table.setStyle(
                    TableStyle(
                        [
                            ("BOX", (0, 0), (-1, -1), 0.5, col(LINE)),
                            ("INNERGRID", (0, 0), (-1, -1), 0.5, col(LINE)),
                            ("VALIGN", (0, 0), (-1, -1), "TOP"),
                            ("TOPPADDING", (0, 0), (-1, -1), 6),
                            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                        ]
                    )
                )
                story.append(table)
            elif isinstance(block, Chart):
                values = [v if v is not None else 0 for v in block.values]
                cats = block.categories
                if block.kind == "bar":
                    h = max(120, 22 * len(cats) + 30)
                    d = Drawing(page_w, h)
                    chart = HorizontalBarChart()
                    chart.x, chart.y, chart.width, chart.height = 110, 15, page_w - 170, h - 25
                    chart.data = [list(reversed(values))]
                    chart.categoryAxis.categoryNames = list(reversed(cats))
                    chart.bars[0].fillColor = col(SERIES[0])
                    chart.bars.strokeColor = None
                    chart.barLabelFormat = number_label
                    chart.barLabels.nudge = 20
                    chart.barLabels.fontSize = 7
                    chart.barLabels.fontName = "DejaVu"
                else:
                    h = 190
                    d = Drawing(page_w, h)
                    if block.kind == "line":
                        chart = HorizontalLineChart()
                        chart.x, chart.y, chart.width, chart.height = 50, 30, page_w - 70, h - 45
                        chart.data = [values]
                        chart.categoryAxis.categoryNames = cats
                        chart.categoryAxis.labels.angle = 45 if len(cats) > 8 else 0
                        chart.categoryAxis.labels.boxAnchor = "ne" if len(cats) > 8 else "n"
                        chart.lines[0].strokeColor = col(SERIES[0])
                        chart.lines[0].strokeWidth = 2
                    else:
                        chart = Pie()
                        chart.x, chart.y, chart.width, chart.height = 40, 10, h - 30, h - 30
                        chart.data = values
                        chart.labels = [f"{c} ({fmt(v)})" for c, v in zip(cats, values, strict=True)]
                        chart.sideLabels = True
                        chart.slices.fontName = "DejaVu"
                        chart.slices.fontSize = 7
                        for i in range(len(values)):
                            chart.slices[i].fillColor = col(SERIES[i % len(SERIES)])
                if block.kind != "pie":
                    # Start at zero (or below for negatives) so differences are not exaggerated.
                    chart.valueAxis.valueMin = min(0.0, min(values, default=0.0))
                    chart.valueAxis.labelTextFormat = lambda v: fmt(v, 0)
                    chart.valueAxis.labels.fontSize = 7
                    chart.valueAxis.labels.fontName = "DejaVu"
                    chart.categoryAxis.labels.fontSize = 7
                    chart.categoryAxis.labels.fontName = "DejaVu"
                    chart.valueAxis.visibleGrid = True
                    chart.valueAxis.gridStrokeColor = col(LINE)
                d.add(chart)
                story.append(KeepTogether([d] + ([p(block.note, "small")] if block.note else [])))
            elif isinstance(block, TableBlock):
                head = [p(c, "head") for c in block.columns]
                body = [[p(v, "cellr" if i in block.numeric else "cell") for i, v in enumerate(row)] for row in block.rows]
                shares = block.widths or [1.0] * len(block.columns)
                widths = [page_w * s / sum(shares) for s in shares]
                table = Table([head, *body], colWidths=widths, repeatRows=1)
                table.setStyle(
                    TableStyle(
                        [
                            ("BACKGROUND", (0, 0), (-1, 0), col(SURFACE_2)),
                            ("LINEBELOW", (0, 0), (-1, -1), 0.4, col(LINE)),
                            ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ]
                    )
                )
                story.append(table)
                if block.note:
                    story.append(p(block.note, "small"))
            elif isinstance(block, Bullets):
                story += [Paragraph("•&nbsp;&nbsp;" + escape(item), styles["body"]) for item in block.items]
            elif isinstance(block, Text):
                story += [p(t, "body") for t in block.paragraphs]
            elif isinstance(block, FindingCard):
                rows = [
                    [p(block.title, "card")],
                    [
                        Paragraph(
                            escape("  ·  ".join(block.badges)), ParagraphStyle("badge", parent=styles["label"], textColor=col(block.colour))
                        )
                    ],
                ]
                rows += [[Table([[p(k, "label")], [p(v, "body")]])] for k, v in block.fields]
                card = Table(rows, colWidths=[page_w])
                card.setStyle(
                    TableStyle(
                        [
                            ("BOX", (0, 0), (-1, -1), 0.6, col(LINE)),
                            ("LINEBEFORE", (0, 0), (0, -1), 3, col(block.colour)),
                            ("LEFTPADDING", (0, 0), (-1, -1), 10),
                        ]
                    )
                )
                story.append(card)
            story.append(Spacer(1, 6))

    def footer(canvas, _doc) -> None:
        canvas.saveState()
        canvas.setFont("DejaVu", 7.5)
        canvas.setFillColor(col(MUTED))
        canvas.drawString(18 * mm, 10 * mm, doc.footer[:150])
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Page {canvas.getPageNumber()}")
        canvas.restoreState()

    buf = io.BytesIO()
    SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=18 * mm,
        title=doc.title,
        author="TabulaVibe AI",
        subject=doc.subtitle,
    ).build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()
